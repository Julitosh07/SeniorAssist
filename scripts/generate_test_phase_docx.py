"""
Script para generar el documento formal de Word (.docx) para la Fase 4: TEST
SeniorAssist_Fase_4_TEST_Informe_Detallado.docx
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "Documentos")
os.makedirs(DOCS_DIR, exist_ok=True)

# Paleta de Colores
COLOR_PRIMARY = RGBColor(26, 54, 93)      # Navy Profundo (#1A365D)
COLOR_SECONDARY = RGBColor(13, 118, 133)  # Teal / Cian (#0D7685)
COLOR_TEXT = RGBColor(45, 55, 72)         # Charcoal Oscuro (#2D3748)
COLOR_MUTED = RGBColor(100, 116, 139)     # Gris Pizarra (#64748B)
COLOR_ACCENT = RGBColor(217, 119, 6)      # Ámbar (#D97706)
COLOR_SUCCESS = RGBColor(5, 150, 105)     # Verde Éxito (#059669)

HEX_PRIMARY = "1A365D"
HEX_SECONDARY = "0D7685"
HEX_LIGHT_BG = "F8FAFC"
HEX_CALLOUT_BG = "F0F9FF"
HEX_CALLOUT_BORDER = "0D7685"
HEX_WARN_BG = "FFFBEB"
HEX_WARN_BORDER = "D97706"
HEX_EVIDENCE_BG = "F1F5F9"
HEX_EVIDENCE_BORDER = "3B82F6"
HEX_BORDER = "CBD5E1"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, hex_color=HEX_BORDER):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_banner(doc, title, subtitle, meta_items):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    meta_table = doc.add_table(rows=len(meta_items), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    for i, (label, val) in enumerate(meta_items):
        row = meta_table.rows[i]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.3)
        cell_val.width = Inches(4.4)
        set_cell_background(cell_lbl, "EDF2F7")
        set_cell_background(cell_val, HEX_LIGHT_BG)
        set_cell_margins(cell_lbl, 60, 60, 100, 100)
        set_cell_margins(cell_val, 60, 60, 100, 100)

        p1 = cell_lbl.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_PRIMARY

        p2 = cell_val.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(val)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_TEXT

    set_table_borders(meta_table, "CBD5E1")
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_h1(doc, text):
    h = doc.add_heading(level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return h

def add_h2(doc, text):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return h

def add_h3(doc, text):
    h = doc.add_heading(level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_TEXT
    return h

def add_body_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_TEXT
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc, title, text, box_type="info"):
    bg_hex = HEX_CALLOUT_BG if box_type == "info" else HEX_WARN_BG
    border_hex = HEX_CALLOUT_BORDER if box_type == "info" else HEX_WARN_BORDER
    title_color = COLOR_SECONDARY if box_type == "info" else COLOR_ACCENT

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.7)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, 120, 120, 180, 180)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_tit = p.add_run(f"📌 {title}\n")
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(10.5)
    r_tit.font.bold = True
    r_tit.font.color.rgb = title_color

    r_txt = p.add_run(text)
    r_txt.font.name = "Calibri"
    r_txt.font.size = Pt(10)
    r_txt.font.italic = True
    r_txt.font.color.rgb = COLOR_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_evidence_box(doc, evidence_num, title, instructions, what_to_show):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.7)
    set_cell_background(cell, HEX_EVIDENCE_BG)
    set_cell_margins(cell, 140, 140, 180, 180)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="{HEX_EVIDENCE_BORDER}"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="{HEX_EVIDENCE_BORDER}"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_EVIDENCE_BORDER}"/>'
        f'  <w:right w:val="single" w:sz="12" w:space="0" w:color="{HEX_EVIDENCE_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r_head = p.add_run(f"📷 ESPACIO PARA CAPTURA DE EVIDENCIA {evidence_num}: {title.upper()}\n")
    r_head.font.name = "Calibri"
    r_head.font.size = Pt(11)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(30, 64, 175) # Azul Real

    p_body = cell.add_paragraph()
    p_body.paragraph_format.space_after = Pt(4)
    p_body.paragraph_format.line_spacing = 1.15
    r_inst = p_body.add_run(f"Cómo obtener esta captura: {instructions}\n\n")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(9.5)
    r_inst.font.italic = True
    r_inst.font.color.rgb = COLOR_TEXT

    r_show_title = p_body.add_run("Qué debe observarse claramente en la imagen:\n")
    r_show_title.font.name = "Calibri"
    r_show_title.font.size = Pt(9.5)
    r_show_title.font.bold = True
    r_show_title.font.color.rgb = COLOR_PRIMARY

    for item in what_to_show:
        p_item = cell.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_item = p_item.add_run(item)
        r_item.font.name = "Calibri"
        r_item.font.size = Pt(9)
        r_item.font.color.rgb = COLOR_TEXT

    # Espacio para pegar la imagen
    p_box = cell.add_paragraph()
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(8)
    r_box = p_box.add_run("\n[ --- PEGAR AQUÍ LA CAPTURA DE PANTALLA --- ]\n")
    r_box.font.name = "Consolas"
    r_box.font.size = Pt(10)
    r_box.font.bold = True
    r_box.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.7)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, 100, 100, 140, 140)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="94A3B8"/>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def generate_test_phase_docx():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    add_header_banner(
        doc,
        title="INFORME TÉCNICO EXHAUSTIVO: FASE 4 - TEST",
        subtitle="Estrategia de Continuous Testing, Calidad de Software y Validación en la Nube",
        meta_items=[
            ("Institución Educativa:", "Unidades Tecnológicas de Santander (UTS)"),
            ("Asignatura / Semestre:", "Nuevas Tecnologías de Desarrollo – 6° Semestre"),
            ("Equipo de Desarrollo:", "SeniorAssist DevOps & QA Engineering Team"),
            ("Fase Abordada:", "FASE 4: TEST (Pruebas Continuas y Calidad de Software)"),
            ("Total de Pruebas Automáticas:", "56 Pruebas (40 Cliente, 7 Voz, 9 Backend) - 100% PASS"),
            ("Servicio en Producción:", "https://seniorassist.onrender.com"),
            ("Consola Interactiva (Swagger UI):", "https://seniorassist.onrender.com/docs"),
            ("Pipeline de CI:", ".github/workflows/test.yml (GitHub Actions)")
        ]
    )

    add_h1(doc, "1. RESUMEN EJECUTIVO DE LA FASE DE TEST")
    add_body_p(doc,
        "En el ciclo de vida DevOps, la fase de TEST (Pruebas Continuas) no es un evento aislado que se ejecuta al finalizar la programación; "
        "es una puerta de calidad automatizada (Quality Gate) que valida de forma continua e inmutable que ningún cambio en el código fuente degrade "
        "la experiencia del usuario, rompa la accesibilidad o comprometa la seguridad criptográfica del sistema."
    )
    add_body_p(doc,
        "Para SeniorAssist, el aseguramiento de calidad presentaba un doble desafío: por un lado, una población vulnerable de adultos mayores "
        "con estrictas demandas de ergonomía cognitiva y visual (WCAG 2.1 AAA, fuentes >= 18sp, botones >= 64dp); y por otro lado, una arquitectura "
        "híbrida distribuida compuesta por un cliente móvil Android con funcionamiento Offline-First (SQLite local) y un microservicio inteligente "
        "en la nube desplegado en Render (FastAPI / Python 3.10)."
    )

    add_h1(doc, "2. MATRIZ GENERAL DE HERRAMIENTAS DE PRUEBA EN SENIORASSIST")

    tools_summary = [
        ("Pytest", "Pruebas unitarias y de integración de backend en Python.", "backend/tests/test_api_suite.py", "9/9 Tests Pasados (1.54s)"),
        ("FastAPI TestClient (HTTPX)", "Simulación de peticiones HTTP en memoria sin sockets TCP.", "backend/tests/test_api_suite.py", "Integrado en Pytest"),
        ("Node.js Test Runner", "Auditoría estricta de accesibilidad WCAG AAA y SQLite.", "scripts/verify_client_integrity.js", "40/40 Tests Pasados (0.3s)"),
        ("Suite de Voz (Node.js)", "Pruebas de lenguaje natural con modismos colombianos.", "scripts/test_voice_suite.js", "7/7 Tests Pasados"),
        ("GitHub Actions (test.yml)", "Orquestación de CI y Quality Gate en cada commit.", ".github/workflows/test.yml", "Pipeline en Verde (Pass)"),
        ("Swagger UI & Render Logs", "Pruebas de humo en vivo y auditoría de peticiones en nube.", "seniorassist.onrender.com/docs", "Producción 200 OK")
    ]

    t_summary = doc.add_table(rows=len(tools_summary) + 1, cols=4)
    t_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_summary.autofit = False
    ts_widths = [Inches(1.8), Inches(2.2), Inches(1.7), Inches(1.0)]

    for j, h_txt in enumerate(["Herramienta", "Alcance Técnico", "Ubicación en el Repo", "Resultado"]):
        c = t_summary.rows[0].cells[j]
        c.width = ts_widths[j]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_txt)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for i, r_vals in enumerate(tools_summary):
        row = t_summary.rows[i + 1]
        bg = HEX_LIGHT_BG if i % 2 == 0 else "FFFFFF"
        for j, txt in enumerate(r_vals):
            c = row.cells[j]
            c.width = ts_widths[j]
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 100, 100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            r.font.size = Pt(8.5)
            r.font.color.rgb = COLOR_PRIMARY if j == 0 else COLOR_TEXT
            if j == 0:
                r.font.bold = True

    set_table_borders(t_summary)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ==========================================
    # HERRAMIENTAS EN DETALLE
    # ==========================================

    # --- HERRAMIENTA 1: PYTEST ---
    add_h1(doc, "3. HERRAMIENTA 1: PYTEST (FRAMEWORK DE PRUEBAS DE BACKEND)")
    add_body_p(doc,
        "Pytest es el marco de trabajo estándar de pruebas unitarias y de integración para el lenguaje Python. Se caracteriza por su "
        "sintaxis minimalista basada en la instrucción nativa 'assert', eliminación de código boilerplate (no requiere heredar de clases "
        "complejas como unittest.TestCase), detección automática de archivos de prueba y fixtures altamente configurables.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Declaración en dependencias: Se configuró en backend/requirements.txt y en el pipeline de CI.\n"
        "2. Creación de la suite: En backend/tests/test_api_suite.py se implementaron 9 funciones de prueba independientes que evalúan "
        "desde el healthcheck hasta la lista blanca estricta de 16 acciones autorizadas y la emisión de tokens JWT HS256.\n"
        "3. Automatización local: Se enlazó a los scripts de npm permitiendo ejecutar 'npm run test:backend' o 'python -m pytest backend/tests/test_api_suite.py'.\n"
        "4. Quality Gate en CI: Integrado en GitHub Actions para bloquear cualquier commit defectuoso.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="1",
        title="Ejecución Exitosa de Pytest en Consola",
        instructions="Abre una terminal en la raíz del proyecto y ejecuta: 'python -m pytest backend/tests/test_api_suite.py'. Toma la captura del resultado.",
        what_to_show=[
            "El comando ejecutado en la consola: 'python -m pytest backend/tests/test_api_suite.py'.",
            "Los 9 puntos verdes que representan cada prueba superada al 100%.",
            "El resumen en verde: '9 passed in 1.54s'.",
            "La versión de Python y Pytest en el encabezado."
        ]
    )

    add_body_p(doc,
        "Beneficios: Sintaxis intuitiva sin clases, auto-descubrimiento de pruebas, ejecución ultra rápida (1.54 segundos para toda la suite) "
        "y reporte minucioso de fallos. Limitaciones: Diseñado exclusivamente para Python; no evalúa código TypeScript ni la interfaz móvil.",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- HERRAMIENTA 2: FASTAPI TESTCLIENT ---
    add_h1(doc, "4. HERRAMIENTA 2: FASTAPI TESTCLIENT BASADO EN HTTPX")
    add_body_p(doc,
        "FastAPI TestClient es una utilidad de pruebas provista por FastAPI y Starlette, construida sobre la librería cliente HTTPX. "
        "Permite simular peticiones HTTP completas contra la aplicación en memoria, sin necesidad de abrir sockets TCP, levantar un servidor ASGI real "
        "ni ocupar puertos de red.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Instanciación en memoria: Se importa 'TestClient' y la instancia 'app' de backend.app.main.\n"
        "2. Simulación de peticiones REST: Ejecuta llamadas GET y POST hacia /health, /auth/device, /ai/intent y /monitoring/metrics.\n"
        "3. Validación de contratos y Pydantic: Evalúa que los datos devueltos coincidan con los esquemas de Pydantic y las cabeceras Bearer.\n"
        "4. Casos negativos: Envía intencionalmente tokens JWT inválidos o corruptos para confirmar que el servidor responda con HTTP 401 Unauthorized.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="2",
        title="Código Fuente de Pruebas con TestClient en VS Code",
        instructions="Abre en Visual Studio Code el archivo 'backend/tests/test_api_suite.py' mostrando las líneas donde se instancia TestClient y se valida /auth/device.",
        what_to_show=[
            "La importación: 'from fastapi.testclient import TestClient'.",
            "La creación del cliente: 'client = TestClient(app)'.",
            "La función 'test_transparent_device_auth()' donde se simula el handshake de dispositivo.",
            "Las aserciones 'assert response.status_code == 200' y 'assert data['token_type'] == 'bearer''."
        ]
    )

    add_body_p(doc,
        "Beneficios: Velocidad extrema en milisegundos, cero conflictos de puertos ocupados en servidores de CI, y fidelidad absoluta a la lógica de middlewares. "
        "Limitaciones: Al ejecutarse en memoria, no evalúa la latencia física de la red ni problemas de certificados SSL reales.",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- HERRAMIENTA 3: NODE.JS TEST RUNNER (CLIENTE Y ACCESIBILIDAD) ---
    add_h1(doc, "5. HERRAMIENTA 3: NODE.JS TEST RUNNER & SUITE DE INTEGRIDAD DEL CLIENTE")
    add_body_p(doc,
        "Es el motor de pruebas automatizado desarrollado específicamente para el cliente móvil en React Native. Utiliza el módulo nativo "
        "'assert' de Node.js y analizadores estáticos de contratos para auditar las directivas de accesibilidad WCAG 2.1 AAA, la norma IEEE 830-1998, "
        "la estructura de base de datos SQLite y los permisos en AndroidManifest.xml.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Creación del script auditor: Se construyó scripts/verify_client_integrity.js con 40 pruebas divididas en 6 secciones temáticas.\n"
        "2. Verificación matemática de ergonomía: Audita que el texto base sea >= 18sp, botones táctiles >= 64dp y colores con contraste >= 7:1.\n"
        "3. Validación Offline-First: Verifica la presencia de las 9 tablas locales en SQLite (contacts, medications, alarms, etc.).\n"
        "4. Integración en npm: Declarado en package.json como 'npm run test:client' y ejecutado en el pipeline de GitHub Actions.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="3",
        title="Auditoría de 40 Pruebas de Integridad de Cliente en Consola",
        instructions="Abre una terminal y ejecuta: 'npm run test:client'. Toma la captura donde se muestre el resumen verde con las 40 pruebas superadas.",
        what_to_show=[
            "El comando ejecutado: 'npm run test:client'.",
            "Las 6 secciones temáticas con sus marcas verdes de verificación.",
            "El resumen final: 'RESULTADO DE LA AUDITORÍA: 40/40 PRUEBAS SUPERADAS'.",
            "El mensaje de éxito: 'TODAS LAS PRUEBAS DE INTEGRIDAD DEL CLIENTE FUERON EXITOSAS'."
        ]
    )

    add_body_p(doc,
        "Beneficios: Cero dependencias pesadas de Babel o emuladores; ejecución en solo 0.3 segundos; garantiza matemáticamente que nadie altere la accesibilidad. "
        "Limitaciones: No renderiza componentes gráficos en pantalla física (eso se valida con el APK instalado en el celular).",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- HERRAMIENTA 4: SUITE DE VOZ ---
    add_h1(doc, "6. HERRAMIENTA 4: SUITE DE PRUEBAS DE LENGUAJE NATURAL Y DIALECTOS DE VOZ")
    add_body_p(doc,
        "Suite de pruebas sintéticas en Node.js (scripts/test_voice_suite.js) que valida la precisión semántica y la tolerancia a modismos "
        "coloquiales colombianos utilizados por adultos mayores al interactuar con el asistente de voz omnipresente.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Definición de dialectos: Pruebas con lenguaje informal ('mama' sin tilde, posesivos 'mi mamá', verbos en infinitivo 'llamar a Carlos').\n"
        "2. Wake-word parsing: Valida la activación por palabra clave ('asistente llama a Pedro').\n"
        "3. Desambiguación de contactos: Comprueba que el sistema distinga estrictamente entre 'Mamá' y 'Mamá Celular' evitando llamadas accidentales.\n"
        "4. Automatización: Enlazado al comando 'npm run test:voice' y al pipeline de CI.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="4",
        title="Ejecución de la Suite de Comandos de Voz y Parsing",
        instructions="Ejecuta en consola: 'npm run test:voice'. Captura la pantalla mostrando las 7 pruebas de voz superadas.",
        what_to_show=[
            "El comando ejecutado: 'npm run test:voice'.",
            "Las 7 líneas con etiqueta [VOZ PASS] demostrando reconocimiento de modismos, tildes y wake-words.",
            "El cuadro de resumen: 'RESULTADO DE VOZ: 7/7 PRUEBAS SUPERADAS'."
        ]
    )

    add_body_p(doc,
        "Beneficios: Previene frustración en adultos mayores asegurando comprensión de habla natural informal. Limitaciones: Evalúa texto transcrito, no ruido acústico del micrófono.",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- ANÁLISIS JEST VS NODE RUNNER ---
    add_h1(doc, "7. ANÁLISIS COMPARATIVO DE HERRAMIENTAS: NODE.JS TEST RUNNER VS. JEST")
    add_body_p(doc,
        "Una de las preguntas más frecuentes de los evaluadores es por qué se optó por un Test Runner optimizado en Node.js en lugar de Jest. "
        "La respuesta radica en la naturaleza de React Native y la filosofía DevOps de eficiencia en CI:"
    )

    jest_comp = [
        ("Criterio Técnico", "Node.js Custom Runner (SeniorAssist)", "Jest Framework"),
        ("Tiempo de Ejecución en CI", "0.3 segundos", "15 a 35 segundos"),
        ("Dependencias Requeridas", "Cero (Nativo de Node.js)", "> 150 MB (Babel, jest-expo, ts-jest)"),
        ("Manejo de Hardware Nativo", "Auditoría directa de contratos y esquemas", "Requiere mockear manualmente cada API nativa"),
        ("Estabilidad ante Actualizaciones", "Muy alta (inmune a cambios de versión)", "Frágil (conflictos ESM vs CJS en React Native)"),
        ("Idoneidad para SeniorAssist", "Óptima para normatividad WCAG e IEEE 830", "Excesivamente pesada para contratos estáticos")
    ]

    t_jest = doc.add_table(rows=len(jest_comp), cols=3)
    t_jest.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_jest.autofit = False
    tj_widths = [Inches(2.1), Inches(2.3), Inches(2.3)]

    for i, r_vals in enumerate(jest_comp):
        row = t_jest.rows[i]
        is_h = (i == 0)
        bg = HEX_PRIMARY if is_h else (HEX_LIGHT_BG if i % 2 == 1 else "FFFFFF")
        for j, txt in enumerate(r_vals):
            c = row.cells[j]
            c.width = tj_widths[j]
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 100, 100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(txt)
            r.font.size = Pt(9)
            if is_h:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                r.font.color.rgb = COLOR_TEXT
                if j == 0:
                    r.font.bold = True

    set_table_borders(t_jest)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- HERRAMIENTA 5: GITHUB ACTIONS ---
    add_h1(doc, "8. HERRAMIENTA 5: GITHUB ACTIONS (TEST.YML) - ORQUESTADOR CI")
    add_body_p(doc,
        "GitHub Actions es la plataforma de integración continua nativa de GitHub. En la fase de TEST, actúa como el 'Quality Gate' "
        "desatendido que ejecuta automáticamente todas las suites de prueba en máquinas virtuales efímeras de Ubuntu cada vez que se sube un commit o PR.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Definición declarativa: En .github/workflows/test.yml se programaron disparadores en push y pull_request hacia main.\n"
        "2. Aprovisionamiento: Descarga una máquina limpia Ubuntu, aprovisiona Node.js 20 y Python 3.10 con caché de paquetes.\n"
        "3. Secuencia de 4 pasos de calidad: (1) Typecheck con TypeScript 5.3, (2) 40 pruebas de cliente, (3) 7 pruebas de voz, (4) 9 pruebas de backend.\n"
        "4. Bloqueo estricto: Si una sola prueba falla, el commit queda en estado rojo ❌ y se prohíbe el merge o despliegue a producción.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="5",
        title="Workflow de Pruebas Exitoso en GitHub Actions",
        instructions="Entra a tu repositorio en GitHub (https://github.com/Julitosh07/SeniorAssist), haz clic en la pestaña 'Actions', selecciona 'DevOps Test Pipeline' y toma la captura del pipeline en verde.",
        what_to_show=[
            "La pestaña 'Actions' del repositorio con el icono verde de verificación (Pass).",
            "El nombre del workflow: 'DevOps Test Pipeline - SeniorAssist'.",
            "Los 4 pasos de prueba ejecutados y completados con éxito.",
            "El tiempo total de ejecución (aproximadamente 1 minuto)."
        ]
    )

    add_body_p(doc,
        "Beneficios: Ejecución en servidores neutros de Ubuntu eliminando sesgos de máquina local; automatización desatendida; historial auditable. "
        "Limitaciones: Requiere conexión a internet y consume minutos de cómputo en la cuota de GitHub.",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- HERRAMIENTA 6: SWAGGER UI & RENDER LOGS ---
    add_h1(doc, "9. HERRAMIENTA 6: SWAGGER UI (/DOCS) & RENDER LIVE LOGS (PRUEBAS EN VIVO EN PRODUCCIÓN)")
    add_body_p(doc,
        "El uso de Swagger UI (OpenAPI 3.0) en https://seniorassist.onrender.com/docs junto a la consola de logs en vivo de Render constituye "
        "una estrategia formal de 'Pruebas de Humo en Producción' (Live Smoke Testing). Permite interactuar directamente contra el microservicio "
        "en la nube y auditar en tiempo real la respuesta en los servidores.",
        bold_prefix="Descripción de la Herramienta: "
    )
    add_body_p(doc,
        "1. Acceso público interactivo: Al ingresar a /docs, la interfaz gráfica permite probar cualquier endpoint sin instalar Postman.\n"
        "2. Prueba de Healthcheck: Ejecución de GET /health para comprobar que el servicio responda HTTP 200 OK en producción.\n"
        "3. Prueba de Autenticación: Ejecución de POST /auth/device enviando el identificador del dispositivo para verificar la firma del token JWT HS256.\n"
        "4. Prueba de NLU: Ejecución de POST /ai/intent con una frase de voz y validación de la respuesta estructurada.\n"
        "5. Auditoría en Render Logs: En la consola web de Render (dashboard.render.com), con el botón 'Live tail' activo, se comprueba la llegada inmediata de cada petición.",
        bold_prefix="Implementación Paso a Paso en DevOps: "
    )

    add_evidence_box(
        doc,
        evidence_num="6A",
        title="Consola Interactiva Swagger UI en Producción Activa",
        instructions="Abre en tu navegador https://seniorassist.onrender.com/docs, despliega el endpoint GET /health o POST /auth/device, dale a 'Try it out' -> 'Execute' y captura la respuesta 200 OK.",
        what_to_show=[
            "La URL en el navegador: https://seniorassist.onrender.com/docs.",
            "El título: 'SeniorAssist Backend - AI & Accessibility Engine'.",
            "El endpoint GET /health o POST /auth/device desplegado con respuesta HTTP 200 OK.",
            "El cuerpo JSON de respuesta devuelto por el servidor en la nube."
        ]
    )

    add_evidence_box(
        doc,
        evidence_num="6B",
        title="Panel de Logs en Vivo en Render con Peticiones Registradas",
        instructions="Entra a dashboard.render.com, entra a 'seniorassist-backend', ve a la pestaña 'Logs', activa 'Live tail' y captura las líneas donde se vean peticiones 200 OK.",
        what_to_show=[
            "El servicio 'seniorassist-backend' en estado verde (Live / Deployed).",
            "La pestaña 'Logs' con el botón 'Live tail' activo.",
            "Líneas de registro de Uvicorn mostrando peticiones HTTP 200 OK:",
            "  INFO: ... - \"GET /health HTTP/1.1\" 200 OK",
            "  INFO: ... - \"POST /auth/device HTTP/1.1\" 200 OK"
        ]
    )

    add_body_p(doc,
        "Beneficios: Demuestra operatividad en el mundo real sobre HTTPS frente a los docentes; no requiere instalar software de terceros; trazabilidad total de IP y códigos HTTP. "
        "Limitaciones: En el plan gratuito de Render, si el servidor pasa más de 15 minutos sin peticiones, el primer arranque experimenta una pausa de 20-30 segundos (cold start).",
        bold_prefix="Beneficios y Limitaciones: "
    )

    # --- DEFENSIVE ARGUMENTATION ---
    add_h1(doc, "10. CÓMO SUSTENTAR LA FASE DE TEST ANTE EL PROFESOR")
    pitch_test = (
        "\"Profesor, en SeniorAssist implementamos una estrategia de Continuous Testing profesional. No nos limitamos a probar que la pantalla se vea bien; "
        "construimos una suite de 56 pruebas automatizadas que corren en cada commit a través de GitHub Actions.\n\n"
        "Para el backend, usamos Pytest y FastAPI TestClient, evaluando no solo casos felices, sino casos de seguridad como el rechazo inmediato de tokens JWT falsos "
        "y el blindaje de la lista blanca de 16 acciones autorizadas.\n\n"
        "Para el cliente móvil, en lugar de usar configuraciones frágiles de Jest que obligan a mockear hardware nativo de forma artificial, desarrollamos un Test Runner "
        "en Node.js que audita matemáticamente que la tipografía cumpla WCAG AAA (>= 18sp), los botones superen los 64dp y las tablas de SQLite garanticen operación Offline-First.\n\n"
        "Y finalmente, para demostrar que el sistema funciona en la nube real, contamos con Swagger UI en Render (/docs), donde podemos lanzar peticiones en vivo "
        "y observar cómo quedan registradas de inmediato en la consola de logs de nuestro servidor desplegado.\""
    )
    add_callout(doc, "Argumentación Maestra de Sustentación para la Fase TEST", pitch_test, box_type="info")

    output_path = os.path.join(DOCS_DIR, "SeniorAssist_Fase_4_TEST_Informe_Detallado.docx")
    doc.save(output_path)
    print(f"[OK] Generado: {output_path}")

if __name__ == "__main__":
    generate_test_phase_docx()
